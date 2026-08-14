from __future__ import annotations

import asyncio
import base64
from io import BytesIO
import json
import os
import socket
import subprocess
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any

import pytest


BACKEND_DIR = Path(__file__).resolve().parents[1]
API_BASE = "http://127.0.0.1:18080"


def request(method: str, path: str, body: dict[str, Any] | None = None) -> Any:
    data = json.dumps(body).encode("utf-8") if body is not None else None
    req = urllib.request.Request(
        API_BASE + path,
        data=data,
        headers={"Content-Type": "application/json", "Connection": "close"},
        method=method,
    )
    with urllib.request.urlopen(req, timeout=30) as response:
        text = response.read().decode("utf-8")
        return json.loads(text) if text else None


def wait_for_server(proc: subprocess.Popen[str]) -> None:
    last_error = "server did not answer"
    for _ in range(60):
        if proc.poll() is not None:
            stdout, stderr = proc.communicate(timeout=2)
            raise RuntimeError(f"Backend exited early. stdout={stdout!r} stderr={stderr!r}")
        try:
            health = request("GET", "/api/health")
            if health.get("status") == "ok":
                return
        except (urllib.error.URLError, TimeoutError, ConnectionError, socket.timeout) as exc:
            last_error = str(exc)
        time.sleep(0.25)
    raise RuntimeError(f"Timed out waiting for backend: {last_error}")


@pytest.fixture(scope="session")
def test_env(tmp_path_factory: pytest.TempPathFactory) -> dict[str, str]:
    root = tmp_path_factory.mktemp("agent_contract")
    workspace = root / "workspace"
    workspace.mkdir()
    (workspace / "sample.py").write_text("def add(a, b):\n    return a + b\n", encoding="utf-8")
    (workspace / "notes.txt").write_text(
        "This workspace is used by pytest to verify context search and summaries.\n",
        encoding="utf-8",
    )

    env = os.environ.copy()
    env.update(
        {
            "AGENT_WORKDIR": str(workspace),
            "HISTORY_DB_PATH": str(root / "sessions-test.db"),
            "ACTIVITY_LOG_PATH": str(root / "activity-test.log"),
            "MYSQL_HOST": "127.0.0.1",
            "MYSQL_PORT": "1",
            "MODEL_NAME": "qwen3:8b",
            "TELEGRAM_BRIDGE_SECRET": "pytest-telegram-secret",
        }
    )
    return env


@pytest.fixture(scope="session")
def backend_server(test_env: dict[str, str]) -> None:
    proc = subprocess.Popen(
        [
            sys.executable,
            "-m",
            "uvicorn",
            "app.main:app",
            "--host",
            "127.0.0.1",
            "--port",
            "18080",
        ],
        cwd=BACKEND_DIR,
        env=test_env,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    try:
        wait_for_server(proc)
        yield
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=10)
        except subprocess.TimeoutExpired:
            proc.kill()
            proc.wait(timeout=5)


def test_email_auth_and_user_session_isolation(backend_server: None) -> None:
    first = request(
        "POST",
        "/api/auth/register",
        {"email": "first@example.com", "password": "secret123", "name": "First User"},
    )
    second = request(
        "POST",
        "/api/auth/register",
        {"email": "second@example.com", "password": "secret123", "name": "Second User"},
    )

    assert first["user"]["email"] == "first@example.com"
    assert second["user"]["email"] == "second@example.com"

    def authed(method: str, path: str, token: str, body: dict[str, Any] | None = None) -> Any:
        data = json.dumps(body).encode("utf-8") if body is not None else None
        req = urllib.request.Request(
            API_BASE + path,
            data=data,
            headers={"Content-Type": "application/json", "Authorization": f"Bearer {token}", "Connection": "close"},
            method=method,
        )
        with urllib.request.urlopen(req, timeout=30) as response:
            text = response.read().decode("utf-8")
            return json.loads(text) if text else None

    session = authed("POST", "/api/sessions", first["token"], {"title": "private first session"})
    first_sessions = authed("GET", "/api/sessions", first["token"])
    second_sessions = authed("GET", "/api/sessions", second["token"])

    assert any(item["id"] == session["id"] for item in first_sessions)
    assert all(item["id"] != session["id"] for item in second_sessions)

    logged_in = request("POST", "/api/auth/login", {"email": "first@example.com", "password": "secret123"})
    assert logged_in["user"]["email"] == "first@example.com"
    logout = authed("POST", "/api/auth/logout", logged_in["token"])
    assert logout == {"ok": True}


def test_auth_rejects_weak_password_and_rate_limits_login(backend_server: None) -> None:
    with pytest.raises(urllib.error.HTTPError) as weak_error:
        request("POST", "/api/auth/register", {"email": "weak@example.com", "password": "12345678", "name": "Weak"})
    weak_error.value.read()
    weak_error.value.close()
    assert weak_error.value.code == 400

    last_code = 0
    for _ in range(9):
        try:
            request("POST", "/api/auth/login", {"email": "limited@example.com", "password": "wrongpass123"})
        except urllib.error.HTTPError as exc:
            last_code = exc.code
            exc.read()
            exc.close()
    assert last_code == 429




def test_telegram_project_isolation(backend_server: None) -> None:
    def telegram_request(method: str, path: str, body: dict[str, Any] | None = None) -> Any:
        data = json.dumps(body).encode("utf-8") if body is not None else None
        req = urllib.request.Request(
            API_BASE + path,
            data=data,
            headers={
                "Content-Type": "application/json",
                "X-Telegram-Bridge-Secret": "pytest-telegram-secret",
                "Connection": "close",
            },
            method=method,
        )
        with urllib.request.urlopen(req, timeout=30) as response:
            text = response.read().decode("utf-8")
            return json.loads(text) if text else None

    first = telegram_request("POST", "/api/telegram/projects", {"telegram_user_id": "111", "title": "Faculty project one"})
    second = telegram_request("POST", "/api/telegram/projects", {"telegram_user_id": "111", "title": "Faculty project two"})
    other = telegram_request("POST", "/api/telegram/projects", {"telegram_user_id": "222", "title": "Other user project"})

    first_user_projects = telegram_request("GET", "/api/telegram/users/111/projects")
    second_user_projects = telegram_request("GET", "/api/telegram/users/222/projects")

    assert {project["id"] for project in first_user_projects} == {first["id"], second["id"]}
    assert {project["id"] for project in second_user_projects} == {other["id"]}

    with pytest.raises(urllib.error.HTTPError) as forbidden_project:
        telegram_request("POST", f"/api/telegram/projects/{first['id']}/message", {"telegram_user_id": "222", "message": "Can I see this?"})
    forbidden_project.value.read()
    forbidden_project.value.close()
    assert forbidden_project.value.code == 404

    with pytest.raises(urllib.error.HTTPError) as bad_secret:
        req = urllib.request.Request(
            API_BASE + "/api/telegram/users/111/projects",
            headers={"X-Telegram-Bridge-Secret": "wrong-secret", "Connection": "close"},
            method="GET",
        )
        urllib.request.urlopen(req, timeout=30)
    bad_secret.value.read()
    bad_secret.value.close()
    assert bad_secret.value.code == 401



def test_telegram_simple_chat_records_history(backend_server: None) -> None:
    def telegram_request(method: str, path: str, body: dict[str, Any] | None = None) -> Any:
        data = json.dumps(body).encode("utf-8") if body is not None else None
        req = urllib.request.Request(
            API_BASE + path,
            data=data,
            headers={
                "Content-Type": "application/json",
                "X-Telegram-Bridge-Secret": "pytest-telegram-secret",
                "Connection": "close",
            },
            method=method,
        )
        with urllib.request.urlopen(req, timeout=30) as response:
            text = response.read().decode("utf-8")
            return json.loads(text) if text else None

    project = telegram_request("POST", "/api/telegram/projects", {"telegram_user_id": "333", "title": "Simple chat"})
    response = telegram_request(
        "POST",
        f"/api/telegram/projects/{project['id']}/message",
        {"telegram_user_id": "333", "message": "Say hello in 2 lines"},
    )
    history = telegram_request("GET", f"/api/telegram/projects/{project['id']}/history?telegram_user_id=333")

    assert "Hello" in response["reply"]
    assert "You: Say hello in 2 lines" in history["summary"]
    assert "SHAMSU:" in history["summary"]
    assert "project_understanding" not in history["summary"]



def test_telegram_code_prompt_creates_workspace_file(backend_server: None, test_env: dict[str, str]) -> None:
    def telegram_request(method: str, path: str, body: dict[str, Any] | None = None) -> Any:
        data = json.dumps(body).encode("utf-8") if body is not None else None
        req = urllib.request.Request(
            API_BASE + path,
            data=data,
            headers={
                "Content-Type": "application/json",
                "X-Telegram-Bridge-Secret": "pytest-telegram-secret",
                "Connection": "close",
            },
            method=method,
        )
        with urllib.request.urlopen(req, timeout=30) as response:
            text = response.read().decode("utf-8")
            return json.loads(text) if text else None

    project = telegram_request("POST", "/api/telegram/projects", {"telegram_user_id": "444", "title": "Code chat"})
    response = telegram_request(
        "POST",
        f"/api/telegram/projects/{project['id']}/message",
        {"telegram_user_id": "444", "message": "Write me a code to print hello world."},
    )

    created = Path(test_env["AGENT_WORKDIR"]) / "telegram_hello_world.py"
    assert created.exists()
    assert 'print("Hello, World!")' in created.read_text(encoding="utf-8")
    assert "telegram_hello_world.py" in response["reply"]
    assert "private SHAMSU Telegram project" not in response["reply"]


def test_telegram_create_file_prompt_writes_safe_file(backend_server: None, test_env: dict[str, str]) -> None:
    def telegram_request(method: str, path: str, body: dict[str, Any] | None = None) -> Any:
        data = json.dumps(body).encode("utf-8") if body is not None else None
        req = urllib.request.Request(
            API_BASE + path,
            data=data,
            headers={
                "Content-Type": "application/json",
                "X-Telegram-Bridge-Secret": "pytest-telegram-secret",
                "Connection": "close",
            },
            method=method,
        )
        with urllib.request.urlopen(req, timeout=30) as response:
            text = response.read().decode("utf-8")
            return json.loads(text) if text else None

    project = telegram_request("POST", "/api/telegram/projects", {"telegram_user_id": "445", "title": "File chat"})
    response = telegram_request(
        "POST",
        f"/api/telegram/projects/{project['id']}/message",
        {"telegram_user_id": "445", "message": "create file telegram_note.txt"},
    )

    created = Path(test_env["AGENT_WORKDIR"]) / "telegram_note.txt"
    assert created.exists()
    assert "Created by SHAMSU Telegram" in created.read_text(encoding="utf-8")
    assert "Created `telegram_note.txt`" in response["reply"]



def test_telegram_natural_code_prompts_follow_user_request(backend_server: None, test_env: dict[str, str]) -> None:
    def telegram_request(method: str, path: str, body: dict[str, Any] | None = None) -> Any:
        data = json.dumps(body).encode("utf-8") if body is not None else None
        req = urllib.request.Request(
            API_BASE + path,
            data=data,
            headers={
                "Content-Type": "application/json",
                "X-Telegram-Bridge-Secret": "pytest-telegram-secret",
                "Connection": "close",
            },
            method=method,
        )
        with urllib.request.urlopen(req, timeout=30) as response:
            text = response.read().decode("utf-8")
            return json.loads(text) if text else None

    project = telegram_request("POST", "/api/telegram/projects", {"telegram_user_id": "446", "title": "Natural code"})
    table = telegram_request(
        "POST",
        f"/api/telegram/projects/{project['id']}/message",
        {"telegram_user_id": "446", "message": "Write me a for loop to print the table of n number"},
    )
    numbers = telegram_request(
        "POST",
        f"/api/telegram/projects/{project['id']}/message",
        {"telegram_user_id": "446", "message": "Write me for loop to print 1 to 100"},
    )
    custom = telegram_request(
        "POST",
        f"/api/telegram/projects/{project['id']}/message",
        {"telegram_user_id": "446", "message": "Can you write a code that will print \"My name is Kakon\""},
    )

    workspace = Path(test_env["AGENT_WORKDIR"])
    assert "telegram_multiplication_table.py" in table["reply"]
    assert "n * i" in (workspace / "telegram_multiplication_table.py").read_text(encoding="utf-8")
    assert "telegram_loop_1_to_100.py" in numbers["reply"]
    assert "range(1, 101)" in (workspace / "telegram_loop_1_to_100.py").read_text(encoding="utf-8")
    assert "My name is Kakon" in custom["reply"]
    assert "My name is Kakon" in (workspace / "telegram_code.py").read_text(encoding="utf-8")
    assert "Hello from SHAMSU Telegram" not in custom["reply"]





def test_telegram_approval_reply_format_includes_yes_no() -> None:
    from app.routes.telegram import _format_approval_request

    reply = _format_approval_request(
        {
            "type": "approval_request",
            "id": "abc123",
            "name": "write_file",
            "path": "app.py",
            "diff": "--- a/app.py\n+++ b/app.py",
            "is_new_file": True,
        }
    )

    assert "Approval required" in reply
    assert "app.py" in reply
    assert "abc123" in reply
    assert "Reply `yes`" in reply
    assert "`no`" in reply


def test_telegram_approval_requires_running_pending_turn(backend_server: None) -> None:
    def telegram_request(method: str, path: str, body: dict[str, Any] | None = None) -> Any:
        data = json.dumps(body).encode("utf-8") if body is not None else None
        req = urllib.request.Request(
            API_BASE + path,
            data=data,
            headers={
                "Content-Type": "application/json",
                "X-Telegram-Bridge-Secret": "pytest-telegram-secret",
                "Connection": "close",
            },
            method=method,
        )
        with urllib.request.urlopen(req, timeout=30) as response:
            text = response.read().decode("utf-8")
            return json.loads(text) if text else None

    project = telegram_request("POST", "/api/telegram/projects", {"telegram_user_id": "448", "title": "Approval route"})
    with pytest.raises(urllib.error.HTTPError) as error:
        telegram_request(
            "POST",
            f"/api/telegram/projects/{project['id']}/approval",
            {"telegram_user_id": "448", "approval_id": "missing", "approved": True},
        )
    error.value.read()
    error.value.close()
    assert error.value.code == 409

def test_telegram_document_upload_indexes_prd_context(backend_server: None, test_env: dict[str, str]) -> None:
    def telegram_request(method: str, path: str, body: dict[str, Any] | None = None) -> Any:
        data = json.dumps(body).encode("utf-8") if body is not None else None
        req = urllib.request.Request(
            API_BASE + path,
            data=data,
            headers={
                "Content-Type": "application/json",
                "X-Telegram-Bridge-Secret": "pytest-telegram-secret",
                "Connection": "close",
            },
            method=method,
        )
        with urllib.request.urlopen(req, timeout=30) as response:
            text = response.read().decode("utf-8")
            return json.loads(text) if text else None

    project = telegram_request("POST", "/api/telegram/projects", {"telegram_user_id": "447", "title": "PRD upload"})
    payload = base64.b64encode(b"OpenBazaar PRD requires buyers, sellers, auctions, COD, and admin audit logs.").decode("ascii")
    uploaded = telegram_request(
        "POST",
        f"/api/telegram/projects/{project['id']}/document",
        {"telegram_user_id": "447", "filename": "OpenBazaar.prd", "data_base64": payload},
    )
    history = telegram_request("GET", f"/api/telegram/projects/{project['id']}/history?telegram_user_id=447")

    workspace = Path(test_env["AGENT_WORKDIR"])
    context_file = workspace / uploaded["path"]
    assert uploaded["kind"] == "text"
    assert uploaded["path"].startswith("uploads/")
    assert context_file.exists()
    assert "buyers, sellers, auctions" in context_file.read_text(encoding="utf-8")
    assert "Uploaded and indexed" in uploaded["reply"]
    assert "Upload: OpenBazaar.prd" in history["summary"]

def test_deployment_profile_uses_safe_env_placeholders() -> None:
    root = Path(__file__).resolve().parents[2]
    compose = (root / "docker-compose.yml").read_text(encoding="utf-8")
    env_example = (root / ".env.deploy.example").read_text(encoding="utf-8")

    assert (root / "backend" / "Dockerfile").exists()
    assert (root / "backend" / ".dockerignore").exists()
    assert (root / "docs" / "DEPLOYMENT.md").exists()
    assert "${SHAMSU_MYSQL_PASSWORD}" in compose
    assert "3307:3306" not in compose
    assert "replace_with_strong_password" in env_example
    assert ".env.deploy" in (root / ".gitignore").read_text(encoding="utf-8")

def test_health_reports_model_and_history_store(backend_server: None) -> None:
    health = request("GET", "/api/health")

    assert health["status"] == "ok"
    assert health["model"]
    assert "sqlite" in health["history_store"] or health["history_store"] == "mysql"
    assert health["workspace"].endswith("workspace")


def test_session_create_state_activity_and_delete(backend_server: None) -> None:
    session = request("POST", "/api/sessions", {"title": "pytest session"})
    session_id = session["id"]

    assert session["title"] == "pytest session"
    assert session["busy"] is False

    state = request("GET", f"/api/sessions/{session_id}/state")
    assert state["events"] == []
    assert state["busy"] is False

    activity = request("GET", f"/api/sessions/{session_id}/activity")
    assert activity["session_id"] == session_id
    assert "prompts" in activity
    assert "tool_calls" in activity

    deleted = request("DELETE", f"/api/sessions/{session_id}")
    assert deleted == {"ok": True}



def test_git_dashboard_read_only_endpoints(backend_server: None) -> None:
    status = request("GET", "/api/git/status")
    assert status["ok"] is True
    assert "branch" in status
    assert "files" in status

    log = request("GET", "/api/git/log?limit=3")
    assert isinstance(log["commits"], list)
    if log["commits"]:
        assert "hash" in log["commits"][0]
        assert "subject" in log["commits"][0]

    search = request("GET", "/api/git/search?query=SHAMSU&limit=5")
    assert search["query"] == "SHAMSU"
    assert isinstance(search["hits"], list)

    diff = request("GET", "/api/git/diff")
    assert "diff" in diff
    assert "truncated" in diff
def test_context_summary_dashboard_search_and_project_map(backend_server: None, test_env: dict[str, str]) -> None:
    workspace = Path(test_env["AGENT_WORKDIR"])
    (workspace / "package.json").write_text('{"dependencies":{"react":"latest","vite":"latest"}}', encoding="utf-8")
    src = workspace / "src"
    src.mkdir(exist_ok=True)
    (src / "main.tsx").write_text('import React from "react";\nimport { App } from "./App";\nexport function boot(){ return App(); }\n', encoding="utf-8")
    (src / "App.tsx").write_text('export function App(){ return <main>hello</main>; }\n', encoding="utf-8")
    summary = request("GET", "/api/context/summary")
    assert "chunk_count" in summary
    assert summary["chunk_count"] >= 1

    dashboard = request("GET", "/api/context/dashboard")
    assert dashboard["file_count"] >= 1
    assert dashboard["chunk_count"] >= 1
    assert dashboard["summary_context_budget"] >= 1
    assert dashboard["conversation_memory_budget"] >= 1
    assert "largest_files" in dashboard
    assert "file_summaries" in dashboard

    overview = request("GET", "/api/context/overview?query=pytest")
    assert overview["query"] == "pytest"
    assert "sample.py" in overview["overview"] or "notes.txt" in overview["overview"]

    search = request("GET", "/api/context/search?query=pytest&limit=3")
    assert search["query"] == "pytest"
    assert isinstance(search["matches"], list)

    (workspace / "auth_notes.txt").write_text(
        "Users sign in with email and password. Login creates a session token for authenticated API calls.\n",
        encoding="utf-8",
    )
    rebuilt = request("POST", "/api/context/vector/rebuild", {"limit_files": 100})
    assert rebuilt["ok"] is True
    assert rebuilt["chunk_count"] >= 1
    assert rebuilt["dims"] == 256

    stats = request("GET", "/api/context/vector/stats")
    assert stats["ready"] is True
    assert stats["chunk_count"] >= rebuilt["chunk_count"]

    vector_search = request("GET", "/api/context/vector/search?query=authentication%20token&limit=3")
    assert vector_search["query"] == "authentication token"
    assert isinstance(vector_search["matches"], list)
    assert any(match["path"] == "auth_notes.txt" for match in vector_search["matches"])


def test_large_file_edit_workflow_plans_ranges_and_verification(backend_server: None, test_env: dict[str, str]) -> None:
    workspace = Path(test_env["AGENT_WORKDIR"])
    large_file = workspace / "large_module.py"
    lines = ["def helper():", "    return 'helper'", ""]
    lines.extend(f"# filler {index}" for index in range(1, 1100))
    lines.extend(["", "def calculate_total(items):", "    return sum(items)"])
    large_file.write_text("\n".join(lines), encoding="utf-8")
    (workspace / "ui.tsx").write_text("export function TotalView(){ return <div>total</div>; }\n", encoding="utf-8")

    result = request(
        "POST",
        "/api/workflows/edit-plan",
        {"prompt": "fix calculate_total in large_module.py and update UI", "query": "calculate_total|TotalView"},
    )

    paths = [item["path"] for item in result["relevant_files"]]
    assert "large_module.py" in paths
    assert any(item["path"] == "large_module.py" and item["large"] for item in result["relevant_files"])
    assert any(window["path"] == "large_module.py" and window["start_line"] < 1110 for window in result["range_windows"])
    assert result["workflow_summary"].startswith("Claude-like large-file workflow")
    assert any("project index" in step.lower() or "project_index" in step.lower() for step in result["claude_like_workflow"])
    assert any("read_file_range" in step for step in result["debug_strategy"])
    assert any(item["path"] == "large_module.py" and item["tool"] == "replace_in_file" for item in result["patch_plan"])
    assert "python -m py_compile <changed_file.py>" in result["verification_commands"]
    assert "npm run build" in result["verification_commands"]
    assert any("bounded range reads" in item for item in result["impact_summary"])



def test_bugfix_and_large_file_plans_are_claude_like() -> None:
    from app.routes.tasks import build_plan

    bugfix = build_plan("fix the traceback in calculator_app.html")
    assert bugfix.mode == "claude-like-bugfix-workflow"
    assert any("index" in step.lower() for step in bugfix.steps)
    assert any("search" in step.lower() for step in bugfix.steps)
    assert any("read_file_range" in item for item in bugfix.stack)
    assert "patch exact block" in bugfix.workflow_summary.lower()

    large = build_plan("fix a bug in a 100000 line huge file")
    assert large.mode == "claude-like-bugfix-workflow" or large.mode == "claude-like-large-file-debugger"
    assert large.requirements_analysis
    assert any("bounded" in item.lower() or "range" in item.lower() for item in large.requirements_analysis)

def test_large_file_edit_workflow_handles_no_match(backend_server: None) -> None:
    result = request("POST", "/api/workflows/edit-plan", {"prompt": "review unknown subsystem", "query": "definitely_missing_symbol"})

    assert result["goal"] == "review unknown subsystem"
    assert "Indexed" in result["index_summary"]
    assert result["next_steps"]


def test_repair_plan_helpers_extract_files_lines_and_query() -> None:
    from app.routes import workflows

    result = workflows.VerificationCommandResult(
        command="python -m pytest -q",
        cwd="workspace",
        ok=False,
        exit_code=1,
        output='File "app/auth.py", line 42, in login_user\nAssertionError: token mismatch',
        failure_summary='python -m pytest -q failed with exit code 1. AssertionError: token mismatch in app/auth.py:42',
    )

    plan = workflows._build_repair_plan([result])

    assert plan[0].likely_files == ["app/auth.py"]
    assert 42 in plan[0].likely_lines
    assert "auth.py" in plan[0].search_query
    assert "next SHAMSU prompt" not in plan[0].next_prompt
    assert "Fix the failure" in plan[0].next_prompt


def test_verification_workflow_dry_run_filters_commands(backend_server: None) -> None:
    result = request(
        "POST",
        "/api/workflows/verify",
        {"target": "workspace", "run": False, "commands": ["python --version", "Remove-Item -Recurse ."]},
    )

    assert result["run"] is False
    assert result["ok"] is False
    assert result["commands"] == ["python --version"]
    assert result["results"] == []
    assert any("Skipped unsupported" in step for step in result["next_steps"])


def test_verification_workflow_runs_and_reports_failure(backend_server: None, test_env: dict[str, str]) -> None:
    workspace = Path(test_env["AGENT_WORKDIR"])
    (workspace / "broken_verify.py").write_text("def broken(:\n    pass\n", encoding="utf-8")

    result = request(
        "POST",
        "/api/workflows/verify",
        {"target": "workspace", "run": True, "commands": ["python -m py_compile broken_verify.py"]},
    )

    assert result["run"] is True
    assert result["ok"] is False
    assert result["results"][0]["ok"] is False
    assert result["repair_feedback"]
    assert "failed" in result["repair_feedback"][0] or "SyntaxError" in result["repair_feedback"][0]


def test_verification_workflow_runs_passing_command(backend_server: None) -> None:
    result = request(
        "POST",
        "/api/workflows/verify",
        {"target": "workspace", "run": True, "commands": ["python --version"]},
    )

    assert result["ok"] is True
    assert result["results"][0]["ok"] is True
    assert result["repair_feedback"] == []

def test_model_list_and_switch_validation(backend_server: None) -> None:
    state = request("GET", "/api/models")
    model_ids = [model["id"] for model in state["models"]]

    assert "qwen3:8b" in model_ids
    assert "qwen3:4b" in model_ids
    assert "qwen3-8k:1.7b" in model_ids

    switched = request("POST", "/api/models/current", {"model_id": "qwen3:4b"})
    assert switched["current"] == "qwen3:4b"

    health = request("GET", "/api/health")
    assert health["model"] == "qwen3:4b"


def test_mcp_tools_list(test_env: dict[str, str]) -> None:
    proc = subprocess.run(
        [sys.executable, "mcp_server.py"],
        input='{"jsonrpc":"2.0","id":1,"method":"tools/list"}\n',
        cwd=BACKEND_DIR,
        env=test_env,
        capture_output=True,
        text=True,
        timeout=60,
    )

    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    tool_names = {tool["name"] for tool in payload["result"]["tools"]}
    assert {"list_directory", "read_file", "search_files", "search_context", "semantic_search", "context_summary", "context_overview", "project_map"}.issubset(tool_names)


def test_cli_sessions_command(backend_server: None, test_env: dict[str, str]) -> None:
    proc = subprocess.run(
        [sys.executable, "cli.py", "--api", API_BASE, "sessions"],
        cwd=BACKEND_DIR,
        env=test_env,
        capture_output=True,
        text=True,
        timeout=60,
    )

    assert proc.returncode == 0, proc.stderr
    assert "No sessions yet." in proc.stdout or "pytest" in proc.stdout or proc.stdout.strip()



def test_cli_direct_file_commands(backend_server: None, test_env: dict[str, str]) -> None:
    write_proc = subprocess.run(
        [sys.executable, "cli.py", "--api", API_BASE, "write", "cli_pytest.txt", "CLI", "direct", "write", "passed"],
        cwd=BACKEND_DIR,
        env=test_env,
        input="y\n",
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert write_proc.returncode == 0, write_proc.stderr
    assert "Approve? [y/N]" in write_proc.stdout
    assert "Wrote" in write_proc.stdout
    assert "[history] recorded in web session" in write_proc.stdout

    sessions_proc = subprocess.run(
        [sys.executable, "cli.py", "--api", API_BASE, "sessions"],
        cwd=BACKEND_DIR,
        env=test_env,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert sessions_proc.returncode == 0, sessions_proc.stderr
    assert "CLI write: cli_pytest.txt" in sessions_proc.stdout

    read_proc = subprocess.run(
        [sys.executable, "cli.py", "--api", API_BASE, "read", "cli_pytest.txt"],
        cwd=BACKEND_DIR,
        env=test_env,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert read_proc.returncode == 0, read_proc.stderr
    assert "CLI direct write passed" in read_proc.stdout

    delete_proc = subprocess.run(
        [sys.executable, "cli.py", "--api", API_BASE, "delete", "cli_pytest.txt"],
        cwd=BACKEND_DIR,
        env=test_env,
        input="y\n",
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert delete_proc.returncode == 0, delete_proc.stderr
    assert "Approve? [y/N]" in delete_proc.stdout
    assert "Deleted cli_pytest.txt" in delete_proc.stdout

    missing_proc = subprocess.run(
        [sys.executable, "cli.py", "--api", API_BASE, "read", "cli_pytest.txt"],
        cwd=BACKEND_DIR,
        env=test_env,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert missing_proc.returncode == 2
    assert "HTTP 404" in missing_proc.stderr


def test_cli_ask_routes_obvious_file_create(backend_server: None, test_env: dict[str, str]) -> None:
    proc = subprocess.run(
        [
            sys.executable,
            "cli.py",
            "--api",
            API_BASE,
            "ask",
            "Create a file named cli_ask_pytest.txt in the workspace with the text: CLI ask route passed.",
        ],
        cwd=BACKEND_DIR,
        env=test_env,
        input="y\n",
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert proc.returncode == 0, proc.stderr
    assert "Approve? [y/N]" in proc.stdout
    assert "Wrote" in proc.stdout

    read_proc = subprocess.run(
        [sys.executable, "cli.py", "--api", API_BASE, "read", "cli_ask_pytest.txt"],
        cwd=BACKEND_DIR,
        env=test_env,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert read_proc.returncode == 0, read_proc.stderr
    assert "CLI ask route passed." in read_proc.stdout



def test_preview_server_start_status_and_stop(test_env: dict[str, str], monkeypatch: pytest.MonkeyPatch) -> None:
    from app.routes import preview

    monkeypatch.setattr(preview.settings, "agent_workdir", test_env["AGENT_WORKDIR"])
    asyncio.run(preview.stop_preview())

    state = asyncio.run(preview.start_preview(preview.PreviewStartRequest(path="sample.py", port=19090)))
    assert state.running is True
    assert state.port == 19090
    assert state.url == "http://127.0.0.1:19090/sample.py"

    status = asyncio.run(preview.preview_status(path="sample.py", port=19090))
    assert status.running is True
    assert status.url == "http://127.0.0.1:19090/sample.py"

    stopped = asyncio.run(preview.stop_preview())
    assert stopped.message in {"Managed preview server stopped.", "No managed preview server was running."}


def test_preview_artifacts_lists_recent_html_outputs(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.routes import preview

    workspace = tmp_path / "workspace"
    app_dir = workspace / "generated_app"
    app_dir.mkdir(parents=True)
    (app_dir / "index.html").write_text("<html><head><title>Generated CRM</title></head><body><h1>Fallback</h1></body></html>", encoding="utf-8")
    (workspace / "notes.txt").write_text("not previewable", encoding="utf-8")
    monkeypatch.setattr(preview.settings, "agent_workdir", str(workspace))

    artifacts = asyncio.run(preview.preview_artifacts(port=19191))

    assert artifacts
    assert artifacts[0].path == "generated_app/index.html"
    assert artifacts[0].title == "Generated CRM"
    assert artifacts[0].url == "http://127.0.0.1:19191/generated_app/index.html"

def _write_and_verify_plan(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, prompt: str) -> tuple[Any, list[str], list[Any]]:
    from app.routes import tasks

    monkeypatch.setattr(tasks.settings, "agent_workdir", str(tmp_path))
    plan = tasks.build_plan(prompt)
    steps: list[tasks.TaskRunStep] = []
    created = tasks._write_suggested_files(plan, overwrite=True, steps=steps)
    ok = tasks._verify_created_files(plan, created, steps)
    assert ok is True
    assert any(step.name == "verify" and step.status == "ok" for step in steps)
    return plan, created, steps


def test_brick_breaker_template_run_creates_previewable_html(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    plan, created, _ = _write_and_verify_plan(tmp_path, monkeypatch, "make a brick breaker game")

    assert plan.mode == "game-generator"
    assert created == ["brick_breaker.html"]
    content = (tmp_path / "brick_breaker.html").read_text(encoding="utf-8")
    assert "Brick Breaker" in content
    assert "paddle" in content
    assert "bricks" in content
    assert "requestAnimationFrame" in content


def test_snake_game_template_run_creates_previewable_html(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    plan, created, _ = _write_and_verify_plan(tmp_path, monkeypatch, "make a snake game in one html file")

    assert plan.mode == "game-generator"
    assert created == ["snake_game.html"]
    content = (tmp_path / "snake_game.html").read_text(encoding="utf-8")
    assert "Snake Game" in content
    assert "ArrowUp" in content
    assert "setInterval" in content


@pytest.mark.parametrize(
    ("prompt", "file_name", "needles"),
    [
        ("make a pong game", "pong.html", ["Pong", "aiPaddle", "requestAnimationFrame"]),
        ("make a tic tac toe game", "tic_tac_toe.html", ["Tic Tac Toe", "checkWinner", "board"]),
        ("make a quiz app", "quiz_app.html", ["Quiz App", "questions", "showQuestion"]),
        ("make a todo app", "todo_app.html", ["Todo App", "localStorage", "renderTodos"]),
        ("make a calculator app", "calculator_app.html", ["Calculator", "appendDigit", "operator"]),
    ],
)
def test_additional_autonomous_templates_run(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, prompt: str, file_name: str, needles: list[str]) -> None:
    _, created, _ = _write_and_verify_plan(tmp_path, monkeypatch, prompt)

    assert created == [file_name]
    content = (tmp_path / file_name).read_text(encoding="utf-8")
    for needle in needles:
        assert needle in content

def test_html_repair_wraps_malformed_generated_content() -> None:
    from app.routes.tasks import _repair_html_content

    repaired = _repair_html_content("demo_game.html", "const score = 0; document.body.textContent = score;")

    assert repaired is not None
    assert "<!doctype html>" in repaired
    assert "<canvas" in repaired
    assert "<script>" in repaired
    assert "const score = 0" in repaired

def test_verification_feedback_reports_broken_script(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.routes import tasks

    monkeypatch.setattr(tasks.settings, "agent_workdir", str(tmp_path))
    (tmp_path / "broken.html").write_text("<html><body><script>function broken(){</script></body></html>", encoding="utf-8")

    feedback = tasks._verification_feedback(["broken.html"])

    assert any("braces look unbalanced" in item for item in feedback)


def test_generic_repair_loop_uses_feedback_to_rewrite_files(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import asyncio
    from app.routes import tasks

    monkeypatch.setattr(tasks.settings, "agent_workdir", str(tmp_path))
    bad_file = {"path": "generic_game.html", "content": "<html><body><script>function broken(){</script></body></html>"}
    fixed_file = {
        "path": "generic_game.html",
        "content": "<!doctype html><html><body><canvas id='game'></canvas><script>function loop(){requestAnimationFrame(loop);} loop();</script></body></html>",
    }

    async def fake_repair(prompt: str, current_files: list[dict[str, str]], feedback: list[str], steps: list[tasks.TaskRunStep], attempt: int) -> list[dict[str, str]]:
        assert prompt == "make a random canvas game"
        assert any("braces look unbalanced" in item for item in feedback)
        steps.append(tasks.TaskRunStep(name="repair-generate", status="ok", detail="fake repair generated one file"))
        return [fixed_file]

    monkeypatch.setattr(tasks, "_generate_json_repair_plan", fake_repair)
    plan = tasks._make_generated_plan("make a random canvas game", [bad_file])
    steps: list[tasks.TaskRunStep] = []
    created = tasks._write_suggested_files(plan, overwrite=True, steps=steps)

    ok, repaired_plan, repaired_files = asyncio.run(tasks._verify_and_repair_loop("make a random canvas game", plan, created, steps))

    assert ok is True
    assert repaired_plan.mode == "json-repaired-task"
    assert repaired_files == ["generic_game.html"]
    assert "requestAnimationFrame" in (tmp_path / "generic_game.html").read_text(encoding="utf-8")
    assert any(step.name == "feedback" for step in steps)
    assert any(step.name == "repair-generate" for step in steps)

def test_crm_management_system_template_run_creates_multi_file_project(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    plan, created, _ = _write_and_verify_plan(tmp_path, monkeypatch, "make a CRM system")

    assert plan.mode == "multi-file-project-generator"
    assert created == [
        "crm_system/package.json",
        "crm_system/README.md",
        "crm_system/WORKFLOW.md",
        "crm_system/PROJECT_MANIFEST.json",
        "crm_system/TASK_BREAKDOWN.md",
        "crm_system/index.html",
        "crm_system/src/main.js",
        "crm_system/src/views.js",
        "crm_system/src/state.js",
        "crm_system/src/data.js",
        "crm_system/src/styles.css",
        "crm_system/tests/smoke_test.py",
    ]
    assert "JavaScript modules" in plan.stack
    assert "crm_system/src/state.js" in plan.file_plan
    assert "multi-file" in plan.notes[0]
    assert "CRM System" in (tmp_path / "crm_system" / "index.html").read_text(encoding="utf-8")
    assert "localStorage" in (tmp_path / "crm_system" / "src" / "state.js").read_text(encoding="utf-8")
    assert "recordsTable" in (tmp_path / "crm_system" / "src" / "views.js").read_text(encoding="utf-8")
    assert "smoke passed" in (tmp_path / "crm_system" / "tests" / "smoke_test.py").read_text(encoding="utf-8")
    manifest = json.loads((tmp_path / "crm_system" / "PROJECT_MANIFEST.json").read_text(encoding="utf-8"))
    task_breakdown = (tmp_path / "crm_system" / "TASK_BREAKDOWN.md").read_text(encoding="utf-8")
    assert manifest["generator"] == "SHAMSU autonomous multi-file project generator"
    assert "Upgrade CRM System to a FastAPI and SQLite backend." in manifest["next_iteration_prompts"]
    assert "Repair Workflow" in task_breakdown


def test_student_management_system_template_run_creates_targeted_multi_file_project(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    plan, created, _ = _write_and_verify_plan(tmp_path, monkeypatch, "build a student management system")

    assert plan.mode == "multi-file-project-generator"
    assert "student_management_system/index.html" in created
    assert "student_management_system/src/data.js" in created
    assert "student_management_system/tests/smoke_test.py" in created
    data = (tmp_path / "student_management_system" / "src" / "data.js").read_text(encoding="utf-8")
    assert "Student Management System" in data
    assert "Program" in data
    assert "Active" in data


def test_website_prompt_creates_previewable_site(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    plan, created, _ = _write_and_verify_plan(tmp_path, monkeypatch, "create a bakery website")

    assert plan.mode == "website-generator"
    assert created == ["bakery_website/index.html", "bakery_website/styles.css", "bakery_website/app.js", "bakery_website/WORKFLOW.md"]
    assert plan.workflow_summary.startswith("prompt -> requirement analysis")
    assert "CSS" in plan.stack
    assert "bakery_website/WORKFLOW.md" in plan.file_plan
    content = (tmp_path / "bakery_website" / "index.html").read_text(encoding="utf-8")
    assert "Bakery Website" in content
    assert "app.js" in content
    workflow = (tmp_path / "bakery_website" / "WORKFLOW.md").read_text(encoding="utf-8")
    assert "Requirement Analysis" in workflow


def test_generic_database_app_prompt_creates_fastapi_sqlite_project(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    plan, created, _ = _write_and_verify_plan(tmp_path, monkeypatch, "build a flashcard tool with database")

    assert plan.mode == "generic-database-app-builder"
    assert "FastAPI" in plan.stack
    assert "SQLite" in plan.stack
    assert "flashcard_tool_database_app/backend/app.py" in created
    assert "flashcard_tool_database_app/schema.sql" in created
    assert "flashcard_tool_database_app/seed.sql" in created
    assert "flashcard_tool_database_app/tests/test_api.py" in created
    assert "CREATE TABLE" in (tmp_path / "flashcard_tool_database_app" / "schema.sql").read_text(encoding="utf-8")
    assert "TestClient" in (tmp_path / "flashcard_tool_database_app" / "tests" / "test_api.py").read_text(encoding="utf-8")
    assert "requirements -> database schema" in plan.workflow_summary

def test_generic_app_prompt_creates_multi_file_project(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    plan, created, _ = _write_and_verify_plan(tmp_path, monkeypatch, "build a flashcard tool")

    assert plan.mode == "generic-multi-file-app-builder"
    assert created == [
        "flashcard_tool/index.html",
        "flashcard_tool/styles.css",
        "flashcard_tool/app.js",
        "flashcard_tool/data.json",
        "flashcard_tool/WORKFLOW.md",
        "flashcard_tool/smoke_test.py",
    ]
    assert plan.workflow_summary == "requirements -> file plan -> create files -> run smoke test -> repair -> preview"
    assert "Open http://127.0.0.1:9000/flashcard_tool/index.html" in plan.verify_commands
    assert "JSON seed data" in plan.stack
    assert "flashcard_tool/smoke_test.py" in plan.file_plan
    html = (tmp_path / "flashcard_tool" / "index.html").read_text(encoding="utf-8")
    app_js = (tmp_path / "flashcard_tool" / "app.js").read_text(encoding="utf-8")
    workflow = (tmp_path / "flashcard_tool" / "WORKFLOW.md").read_text(encoding="utf-8")
    assert "Flashcard Tool" in html
    assert "localStorage" in app_js
    assert "Requirement analysis" in workflow

def test_generic_database_app_prompt_creates_fastapi_sqlite_project(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    plan, created, _ = _write_and_verify_plan(tmp_path, monkeypatch, "build a flashcard tool with database")

    assert plan.mode == "generic-database-app-builder"
    assert "FastAPI" in plan.stack
    assert "SQLite" in plan.stack
    assert "flashcard_tool_database_app/backend/app.py" in created
    assert "flashcard_tool_database_app/schema.sql" in created
    assert "flashcard_tool_database_app/seed.sql" in created
    assert "flashcard_tool_database_app/tests/test_api.py" in created
    assert "CREATE TABLE" in (tmp_path / "flashcard_tool_database_app" / "schema.sql").read_text(encoding="utf-8")
    assert "TestClient" in (tmp_path / "flashcard_tool_database_app" / "tests" / "test_api.py").read_text(encoding="utf-8")
    assert "requirements -> database schema" in plan.workflow_summary

def test_generic_app_prompt_creates_multi_file_project(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    plan, created, _ = _write_and_verify_plan(tmp_path, monkeypatch, "build a flashcard tool")

    assert plan.mode == "generic-multi-file-app-builder"
    assert created == [
        "flashcard_tool/index.html",
        "flashcard_tool/styles.css",
        "flashcard_tool/app.js",
        "flashcard_tool/data.json",
        "flashcard_tool/WORKFLOW.md",
        "flashcard_tool/smoke_test.py",
    ]
    assert plan.workflow_summary == "requirements -> file plan -> create files -> run smoke test -> repair -> preview"
    assert "Open http://127.0.0.1:9000/flashcard_tool/index.html" in plan.verify_commands
    assert "JSON seed data" in plan.stack
    assert "flashcard_tool/smoke_test.py" in plan.file_plan
    html = (tmp_path / "flashcard_tool" / "index.html").read_text(encoding="utf-8")
    app_js = (tmp_path / "flashcard_tool" / "app.js").read_text(encoding="utf-8")
    workflow = (tmp_path / "flashcard_tool" / "WORKFLOW.md").read_text(encoding="utf-8")
    assert "Flashcard Tool" in html
    assert "localStorage" in app_js
    assert "Requirement analysis" in workflow

def test_general_system_prompt_creates_dashboard_prototype(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    plan, created, _ = _write_and_verify_plan(tmp_path, monkeypatch, "make a clinic system")

    assert plan.mode == "system-prototype-generator"
    assert created == ["clinic_system/index.html", "clinic_system/styles.css", "clinic_system/app.js", "clinic_system/WORKFLOW.md"]
    assert "localStorage" in plan.stack
    assert plan.clarification_questions
    content = (tmp_path / "clinic_system" / "index.html").read_text(encoding="utf-8")
    assert "Clinic System" in content
    app_js = (tmp_path / "clinic_system" / "app.js").read_text(encoding="utf-8")
    assert "addRecord" in app_js
    assert "toggleRecord" in app_js
    assert "localStorage" in app_js


def test_advisory_build_plan_explains_workflow_before_result() -> None:
    from app.routes import tasks

    plan = tasks._with_advisory(tasks.build_plan("make a student management website"), "make a student management website")

    assert plan.requirements_analysis
    assert any("dashboard" in item.lower() or "records" in item.lower() for item in plan.requirements_analysis)
    assert plan.clarification_questions
    assert plan.stack
    assert plan.file_plan
    assert "Claude-style build" in plan.workflow_summary

def test_general_planner_routes_unknown_build_to_multi_file_fallback() -> None:
    from app.routes.tasks import build_plan

    plan = build_plan("make a maze game")
    assert plan.mode == "generic-multi-file-app-builder"
    assert len(plan.suggested_files) == 6
    assert "requirements -> file plan" in plan.workflow_summary
    assert any(item["path"].endswith("/smoke_test.py") for item in plan.suggested_files)


def test_admin_dashboard_groups_user_generated_files_without_upload_leakage() -> None:
    from app.routes.admin import _local_projects_from_files, _overall_verification_status

    files = [
        {"path": "crm_system/index.html", "session_id": "s1", "session_title": "CRM"},
        {"path": "crm_system/tests/smoke_test.py", "session_id": "s1", "session_title": "CRM"},
        {"path": "uploads/private.pdf", "session_id": "s1", "session_title": "CRM"},
    ]
    projects = _local_projects_from_files(files)

    assert len(projects) == 1
    assert projects[0]["path"] == "crm_system"
    assert projects[0]["previewable"] is True
    assert projects[0]["verification_status"] == "has smoke test"
    assert "uploads/private.pdf" not in projects[0]["generated_files"]
    assert _overall_verification_status([{"verification_status": "verified"}, {"verification_status": "needs repair"}]) == {"verified": 1, "needs repair": 1}

def test_generated_file_validation_rejects_unsafe_paths() -> None:
    from app.routes.tasks import _validated_generated_files, TaskRunStep

    steps: list[TaskRunStep] = []
    files = _validated_generated_files(
        {
            "files": [
                {"path": "snake.html", "content": "<html><script></script></html>"},
                {"path": "../escape.py", "content": "print('bad')"},
                {"path": "binary.exe", "content": "bad"},
            ]
        },
        steps,
    )

    assert files == [{"path": "snake.html", "content": "<html><script></script></html>"}]
    assert any("outside the workspace" in step.detail or "unsupported" in step.detail for step in steps)

def test_autonomous_task_template_writes_and_verifies_game(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.routes import tasks

    monkeypatch.setattr(tasks.settings, "agent_workdir", str(tmp_path))
    plan = tasks.build_plan("make a bouncing ball game")
    steps: list[tasks.TaskRunStep] = []
    created = tasks._write_suggested_files(plan, overwrite=True, steps=steps)
    ok = tasks._verify_created_files(plan, created, steps)

    assert ok is True
    assert plan.mode == "game-generator"
    assert created == ["bouncing_ball.html"]
    target = tmp_path / "bouncing_ball.html"
    assert target.exists()
    content = target.read_text(encoding="utf-8")
    assert "<canvas" in content
    assert "requestAnimationFrame" in content
    assert any(step.name == "verify" and step.status == "ok" for step in steps)

def test_autonomous_task_returns_reliability_report(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.routes import tasks

    monkeypatch.setattr(tasks.settings, "agent_workdir", str(tmp_path))
    plan = tasks.build_plan("make a bouncing ball game")
    steps: list[tasks.TaskRunStep] = [tasks.TaskRunStep(name="plan", status="ok", detail="Selected game-generator workflow.")]
    created = tasks._write_suggested_files(plan, overwrite=True, steps=steps)
    ok = tasks._verify_created_files(plan, created, steps)

    report = tasks._build_reliability_report(steps, created, ok)

    assert report.final_status == "verified"
    assert report.repair_attempts == 0
    assert any(phase.startswith("plan: ok") for phase in report.phases)
    assert any(phase.startswith("write: ok") for phase in report.phases)
    assert any(phase.startswith("verify: ok") for phase in report.phases)
    assert "Open the preview" in report.next_action
def test_task_plan_returns_bouncing_ball_template() -> None:
    from app.routes.tasks import build_plan

    plan = build_plan("make a bouncing ball game")
    assert plan.mode == "game-generator"
    assert plan.suggested_files[0]["path"] == "bouncing_ball.html"
    assert "requestAnimationFrame" in plan.suggested_files[0]["content"]
def test_cli_index_range_and_patch_commands(backend_server: None, test_env: dict[str, str]) -> None:
    workspace = Path(test_env["AGENT_WORKDIR"])
    patch_target = workspace / "patch_target.py"
    patch_target.write_text("def greet():\n    return 'hello'\n", encoding="utf-8")

    index_proc = subprocess.run(
        [sys.executable, "cli.py", "--api", API_BASE, "index"],
        cwd=BACKEND_DIR,
        env=test_env,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert index_proc.returncode == 0, index_proc.stderr
    assert "sample.py" in index_proc.stdout
    assert "patch_target.py" in index_proc.stdout
    assert "[history] recorded in web session" in index_proc.stdout or "[history warning]" in index_proc.stderr

    range_proc = subprocess.run(
        [sys.executable, "cli.py", "--api", API_BASE, "range", "sample.py", "1", "2"],
        cwd=BACKEND_DIR,
        env=test_env,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert range_proc.returncode == 0, range_proc.stderr
    assert "1: def add" in range_proc.stdout
    assert "2:     return a + b" in range_proc.stdout

    patch_proc = subprocess.run(
        [sys.executable, "cli.py", "--api", API_BASE, "patch", "patch_target.py", "return 'hello'", "return 'hi'"],
        cwd=BACKEND_DIR,
        env=test_env,
        input="y\n",
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert patch_proc.returncode == 0, patch_proc.stderr
    assert "Approve? [y/N]" in patch_proc.stdout
    assert "Patched 'patch_target.py'" in patch_proc.stdout
    assert "return 'hi'" in patch_target.read_text(encoding="utf-8")

def test_implicit_code_fence_path_extraction() -> None:
    from app.agent.loop import _extract_candidate_file_path

    assert _extract_candidate_file_path("**File: `division.py`**") == "division.py"
    assert _extract_candidate_file_path("", "create a new file called game.py") == "game.py"



def test_upload_returns_metadata_and_writes_context_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import asyncio
    from io import BytesIO
    from starlette.datastructures import UploadFile
    from app.routes import uploads

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    monkeypatch.setattr(uploads.settings, "agent_workdir", str(workspace))
    upload = UploadFile(filename="Faculty Notes.txt", file=BytesIO(b"This document explains grading rubric and demo requirements."))

    result = asyncio.run(uploads.upload_context_file(upload))

    assert result["name"] == "Faculty Notes.txt"
    assert result["kind"] == "text"
    assert result["extension"] == "txt"
    assert result["words"] == 8
    assert result["lines"] == 1
    assert result["bytes"] == 60
    assert result["suggested_prompts"][0].startswith("Summarize this uploaded file")
    assert "grading rubric" in result["summary"]
    assert "demo requirements" in result["preview"]
    stored = workspace / str(result["path"])
    assert stored.exists()
    assert "Uploaded source: Faculty Notes.txt" in stored.read_text(encoding="utf-8")




def test_upload_prd_returns_build_mode_analysis(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from fastapi import UploadFile
    from app.routes import uploads

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    monkeypatch.setattr(uploads.settings, "agent_workdir", str(workspace))
    text = """OpenBazaar PRD requirements
Roles: Guest, Buyer, Seller, Admin.
Entities: users, categories, items, bids, orders.
Workflows: register, login, search, auction, checkout, cash on delivery.
Pages: homepage, product page, seller dashboard, buyer dashboard, admin.
Security: password hashing, OTP, rate limit, role-based access.
"""
    result = asyncio.run(uploads.upload_context_file(UploadFile(filename="OpenBazaar_PRD.txt", file=BytesIO(text.encode("utf-8")))))

    assert result["prd_analysis"]["is_prd"] is True
    assert "Do you want me to build Step 1" in result["prd_analysis"]["message"]
    assert "buyer" in result["prd_analysis"]["found"]["roles"]
    assert result["suggested_prompts"][1].startswith("Build Step 1")

def test_upload_accepts_docx_and_extracts_paragraphs(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    pytest.importorskip("docx")
    from docx import Document
    from io import BytesIO
    from starlette.datastructures import UploadFile
    from app.routes import uploads

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    monkeypatch.setattr(uploads.settings, "agent_workdir", str(workspace))
    document = Document()
    document.add_paragraph("OpenBazaar requires buyers, sellers, bids, and COD orders.")
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    result = asyncio.run(uploads.upload_context_file(UploadFile(filename="OpenBazaar PRD.docx", file=buffer)))

    assert result["kind"] == "docx"
    assert result["extension"] == "docx"
    assert "Word document" in result["suggested_prompts"][0]
    assert "COD orders" in result["preview"]

def test_attached_upload_context_is_prioritized_over_workspace_map(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.agent.loop import AgentSession
    from app.agent import loop

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    upload_dir = workspace / "uploads"
    upload_dir.mkdir()
    uploaded = upload_dir / "20260802-120000-abcd-Faculty-Notes.txt"
    uploaded.write_text("Uploaded source: Faculty Notes.txt\n\nRubric says demo must explain upload handling.", encoding="utf-8")
    (workspace / "main.py").write_text("print('workspace file should not replace upload')\n", encoding="utf-8")
    monkeypatch.setattr(loop.settings, "agent_workdir", str(workspace))

    session = AgentSession(title="upload context test")
    packet = session._with_file_context(
        "what does this uploaded file say?",
        ["uploads/20260802-120000-abcd-Faculty-Notes.txt"],
        {"uploads/20260802-120000-abcd-Faculty-Notes.txt": "Faculty Notes.txt"},
    )

    assert "Uploaded attachment: Faculty Notes.txt" in packet
    assert "Rubric says demo must explain upload handling" in packet
    assert "Compact project architecture map" not in packet

def test_implicit_code_fence_becomes_approval(monkeypatch: pytest.MonkeyPatch) -> None:
    import asyncio
    from app.agent.loop import AgentSession

    session = AgentSession(title="implicit edit test")
    session._last_user_message = "create a new file called division.py which can divide numbers"
    content = """Sure.\n\n**File: `division.py`**\n\n```python\ndef divide(a, b):\n    return a / b\n```\n"""
    captured: dict[str, object] = {}

    async def fake_execute_tool(name: str, args: dict[str, object]) -> str:
        captured["name"] = name
        captured["args"] = args
        return "ok"

    monkeypatch.setattr(session, "_execute_tool", fake_execute_tool)
    handled = asyncio.run(session._maybe_offer_implicit_edit(content))

    assert handled is True
    assert captured["name"] == "write_file"
    assert captured["args"] == {"path": "division.py", "content": "def divide(a, b):\n    return a / b\n"}



















def test_web_search_parser_extracts_sourced_results() -> None:
    from app.web_search import format_search_results, parse_duckduckgo_html

    html = """
    <html><body>
      <a class="result__a" href="/l/?kh=-1&amp;uddg=https%3A%2F%2Fexample.com%2Fdocs">Example Docs</a>
      <a class="result__snippet">Useful docs snippet for SHAMSU.</a>
      <a class="result__a" href="https://example.org/news">Example News</a>
      <div class="result__snippet">Current information snippet.</div>
    </body></html>
    """

    results = parse_duckduckgo_html(html, limit=5)
    assert len(results) == 2
    assert results[0].title == "Example Docs"
    assert results[0].url == "https://example.com/docs"
    assert results[0].snippet == "Useful docs snippet for SHAMSU."

    formatted = format_search_results({"ok": True, "query": "shamsu", "results": [result.__dict__ for result in results]})
    assert "Web search results for: shamsu" in formatted
    assert "https://example.com/docs" in formatted


def test_web_search_empty_query_returns_error() -> None:
    from app.web_search import search_web

    result = search_web("   ")
    assert result["ok"] is False
    assert result["error"] == "Query is required."


def test_web_search_route_rejects_missing_query() -> None:
    from fastapi.testclient import TestClient

    from app.main import app

    response = TestClient(app).get("/api/web/search")
    assert response.status_code == 422
def test_agent_and_mcp_expose_web_search_tool(test_env: dict[str, str]) -> None:
    from app.agent.tools import READ_ONLY_TOOLS, TOOL_NAMES, TOOL_SCHEMAS

    assert "web_search" in READ_ONLY_TOOLS
    assert "web_search" in TOOL_NAMES
    assert any(schema["function"]["name"] == "web_search" for schema in TOOL_SCHEMAS)

    proc = subprocess.run(
        [sys.executable, "mcp_server.py"],
        input='{"jsonrpc":"2.0","id":1,"method":"tools/list"}\n',
        cwd=BACKEND_DIR,
        env=test_env,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    tool_names = {tool["name"] for tool in payload["result"]["tools"]}
    assert "web_search" in tool_names







def test_database_backed_crm_template_creates_fastapi_sqlite_project(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    plan, created, _ = _write_and_verify_plan(tmp_path, monkeypatch, "make a CRM system with database backend")

    assert plan.mode == "database-backed-system-generator"
    assert "FastAPI" in plan.stack
    assert "SQLite" in plan.stack
    assert created == [
        "crm_system_database/package.json",
        "crm_system_database/requirements.txt",
        "crm_system_database/README.md",
        "crm_system_database/WORKFLOW.md",
        "crm_system_database/backend/__init__.py",
        "crm_system_database/backend/database.py",
        "crm_system_database/backend/main.py",
        "crm_system_database/frontend/index.html",
        "crm_system_database/frontend/app.js",
        "crm_system_database/frontend/styles.css",
        "crm_system_database/tests/smoke_test.py",
    ]
    backend_main = (tmp_path / "crm_system_database" / "backend" / "main.py").read_text(encoding="utf-8")
    database_py = (tmp_path / "crm_system_database" / "backend" / "database.py").read_text(encoding="utf-8")
    frontend_js = (tmp_path / "crm_system_database" / "frontend" / "app.js").read_text(encoding="utf-8")
    readme = (tmp_path / "crm_system_database" / "README.md").read_text(encoding="utf-8")

    assert "FastAPI" in backend_main
    assert "StaticFiles" in backend_main
    assert "sqlite3" in database_py
    assert "CREATE TABLE IF NOT EXISTS records" in database_py
    assert "/api/records" in frontend_js
    assert "python -m uvicorn backend.main:app" in readme


def test_plain_crm_prompt_still_uses_static_multi_file_project() -> None:
    from app.routes.tasks import build_plan

    plan = build_plan("make a CRM system")
    assert plan.mode == "multi-file-project-generator"
    assert all("database.py" not in item["path"] for item in plan.suggested_files)


def test_full_stack_prompt_creates_fastapi_sqlite_frontend_project(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    plan, created, _ = _write_and_verify_plan(tmp_path, monkeypatch, "make a full-stack clinic app")

    assert plan.mode == "full-stack-project-generator"
    assert "FastAPI" in plan.stack
    assert "SQLite" in plan.stack
    assert created == [
        "clinic_project/package.json",
        "clinic_project/requirements.txt",
        "clinic_project/README.md",
        "clinic_project/WORKFLOW.md",
        "clinic_project/backend/__init__.py",
        "clinic_project/backend/database.py",
        "clinic_project/backend/main.py",
        "clinic_project/frontend/index.html",
        "clinic_project/frontend/app.js",
        "clinic_project/frontend/styles.css",
        "clinic_project/tests/smoke_test.py",
    ]
    backend_main = (tmp_path / "clinic_project" / "backend" / "main.py").read_text(encoding="utf-8")
    database_py = (tmp_path / "clinic_project" / "backend" / "database.py").read_text(encoding="utf-8")
    frontend_js = (tmp_path / "clinic_project" / "frontend" / "app.js").read_text(encoding="utf-8")
    workflow = (tmp_path / "clinic_project" / "WORKFLOW.md").read_text(encoding="utf-8")

    assert "FastAPI" in backend_main
    assert "GET /api/health" not in backend_main
    assert "sqlite3" in database_py
    assert "CREATE TABLE IF NOT EXISTS items" in database_py
    assert "/api/items" in frontend_js
    assert "full-stack app" in workflow
    assert any("uvicorn" in command for command in plan.verify_commands)


def test_non_full_stack_system_still_uses_prototype_generator() -> None:
    from app.routes.tasks import build_plan

    plan = build_plan("make a clinic system")
    assert plan.mode == "system-prototype-generator"


def test_connector_marketplace_lists_enabled_and_planned_tools() -> None:
    from app.routes.connectors import list_connectors

    result = list_connectors()

    connectors = {connector["id"]: connector for connector in result["connectors"]}
    assert result["enabled_count"] >= 5
    assert result["planned_count"] >= 5
    assert connectors["workspace-files"]["status"] == "enabled"
    assert connectors["web-search"]["privacy"] == "external-request"
    assert "web_search" in connectors["web-search"]["tools"]
    assert connectors["google-drive"]["setup_required"] is True
    assert connectors["github-remote-search"]["privacy"] == "token-required"
    assert "version lookup" in connectors["package-registry-search"]["capabilities"]
    assert connectors["deployment-hosting"]["setup_required"] is True

def test_connector_marketplace_get_single_connector() -> None:
    from app.routes.connectors import get_connector

    connector = get_connector("mcp-server")

    assert connector["id"] == "mcp-server"
    assert connector["status"] == "enabled"
    assert "tools/list" in connector["capabilities"]



def test_agent_tools_expose_semantic_search() -> None:
    from app.agent import tools

    names = {schema["function"]["name"] for schema in tools.TOOL_SCHEMAS}
    assert "semantic_search" in names
    assert "semantic_search" in tools.READ_ONLY_TOOLS


def test_vector_index_direct_semantic_search(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import vector_index

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    monkeypatch.setattr(vector_index.settings, "agent_workdir", str(workspace))
    monkeypatch.setattr(vector_index.settings, "history_db_path", str(tmp_path / "sessions.db"))
    (workspace / "auth.py").write_text("def login_user(email, password):\n    return create_session_token(email)\n", encoding="utf-8")
    (workspace / "billing.py").write_text("def create_invoice(total):\n    return total\n", encoding="utf-8")

    rebuilt = vector_index.rebuild_index()
    matches = vector_index.semantic_search("authentication session token", limit=3)

    assert rebuilt["chunk_count"] == 2
    assert matches
    assert matches[0]["path"] == "auth.py"
    assert "login_user" in matches[0]["preview"]



def test_vector_index_uses_path_and_identifier_terms(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import vector_index

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    nested = workspace / "frontend" / "components"
    nested.mkdir(parents=True)
    monkeypatch.setattr(vector_index.settings, "agent_workdir", str(workspace))
    monkeypatch.setattr(vector_index.settings, "history_db_path", str(tmp_path / "sessions.db"))
    (nested / "AccountDashboard.tsx").write_text("export function AccountDashboard(){ return loadUserSessionState(); }\n", encoding="utf-8")
    (workspace / "notes.txt").write_text("plain unrelated notes about invoices\n", encoding="utf-8")

    rebuilt = vector_index.rebuild_index()
    matches = vector_index.semantic_search("account dashboard session state component", limit=3)

    assert rebuilt["dims"] == 256
    assert matches
    assert matches[0]["path"] == "frontend/components/AccountDashboard.tsx"
def test_query_policy_routes_general_web_workspace_and_upload() -> None:
    from app import query_policy

    general = query_policy.classify_query("explain recursion simply", [])
    assert general.route == "general_model"
    assert general.enable_tools is False

    web = query_policy.classify_query("search latest React version online", [])
    assert web.route == "web_search"
    assert web.use_web_search is True
    assert web.enable_tools is True

    workspace = query_policy.classify_query("search this project for auth bug", [])
    assert workspace.route == "workspace_context"
    assert workspace.use_workspace_context is True

    upload = query_policy.classify_query("what does this uploaded file say?", ["uploads/demo.txt"])
    assert upload.route == "uploaded_context"
    assert upload.use_uploaded_context is True


def test_agent_context_injects_web_search_for_current_questions(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.agent.loop import AgentSession
    from app.agent import loop

    monkeypatch.setattr(loop.tools, "web_search", lambda query: "Web search results for: " + query + "\n1. Source\n   https://example.com/current")
    monkeypatch.setattr(loop.context_index, "conversation_memory", lambda events, query: "")
    monkeypatch.setattr(loop.context_index, "automatic_summary_context", lambda query: "SHOULD_NOT_USE_PROJECT_SUMMARY")
    monkeypatch.setattr(loop.context_index, "project_map_context", lambda query: "SHOULD_NOT_USE_PROJECT_MAP")
    monkeypatch.setattr(loop.context_index, "automatic_context", lambda query: "SHOULD_NOT_USE_WORKSPACE_CONTEXT")

    packet = AgentSession(title="web policy test")._with_file_context("what is the latest Python release?", [])

    assert "Context routing decision: web_search" in packet
    assert "Public web search context" in packet
    assert "https://example.com/current" in packet
    assert "SHOULD_NOT_USE_PROJECT" not in packet
    assert "SHOULD_NOT_USE_WORKSPACE" not in packet


def test_agent_context_uses_workspace_for_project_questions(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.agent.loop import AgentSession
    from app.agent import loop

    monkeypatch.setattr(loop.tools, "web_search", lambda query: "SHOULD_NOT_SEARCH_WEB")
    monkeypatch.setattr(loop.context_index, "conversation_memory", lambda events, query: "")
    monkeypatch.setattr(loop.context_index, "automatic_summary_context", lambda query: "summary context")
    monkeypatch.setattr(loop.context_index, "project_map_context", lambda query: "project map context")
    monkeypatch.setattr(loop.context_index, "automatic_context", lambda query: "exact workspace context")
    monkeypatch.setattr(loop.tools, "long_context_bundle", lambda query: "Project memory\n- auth files")

    packet = AgentSession(title="workspace policy test")._with_file_context("search this project for auth bug", [])

    assert "Context routing decision: workspace_context" in packet
    assert "Fused long-context bundle" in packet
    assert "Project memory" in packet
    assert "exact workspace context" in packet
    assert "SHOULD_NOT_SEARCH_WEB" not in packet


def test_general_explanation_does_not_enable_tools() -> None:
    from app.agent.loop import _should_enable_tools

    assert _should_enable_tools("explain photosynthesis simply", []) is False
    assert _should_enable_tools("what is the latest Ollama version?", []) is True
    assert _should_enable_tools("explain this project structure", []) is True


def test_fallback_parser_recovers_function_style_tool_calls() -> None:
    from app.agent.loop import _extract_fallback_tool_call

    shell = _extract_fallback_tool_call('I will run run_shell(command="python app.py") now')
    assert shell == {"name": "run_shell", "arguments": {"command": "python app.py"}}

    read = _extract_fallback_tool_call("Next call read_file(path='main.py') before editing")
    assert read == {"name": "read_file", "arguments": {"path": "main.py"}}

    wrapped = _extract_fallback_tool_call('{"function":{"name":"search_files","arguments":{"query":"TODO"}}}')
    assert wrapped == {"name": "search_files", "arguments": {"query": "TODO"}}


def test_tool_loop_marks_verification_after_mutating_tools() -> None:
    from app.agent.loop import AgentSession

    session = AgentSession(title="tool loop state test")
    session._update_tool_loop_state("write_file", "Wrote main.py", True)
    assert session._verification_pending is True
    assert session._should_continue_for_verification("Done, I created the file.") is True

    session._update_tool_loop_state("run_shell", "[exit code: 1]", True)
    assert session._verification_pending is True
    assert session._repair_pending is True

    session._update_tool_loop_state("run_shell", "all good\n[exit code: 0]", True)
    assert session._verification_pending is False
    assert session._repair_pending is False


def test_tool_loop_supervisor_continues_until_verification(monkeypatch: pytest.MonkeyPatch) -> None:
    import asyncio
    from app.agent.loop import AgentSession

    session = AgentSession(title="tool loop supervisor test")
    session._verification_pending = True
    responses = iter([
        ("I created the file.", []),
        ("", [{"name": "run_shell", "arguments": {"command": "python --version"}}]),
        ("Verified successfully.", []),
    ])
    executed: list[tuple[str, dict[str, object]]] = []

    async def fake_get_model_response() -> tuple[str, list[dict[str, object]]]:
        return next(responses)

    async def fake_execute_tool(name: str, args: dict[str, object]) -> str:
        executed.append((name, args))
        session._update_tool_loop_state(name, "Python 3.13\n[exit code: 0]", True)
        return "Python 3.13\n[exit code: 0]"

    monkeypatch.setattr(session, "_get_model_response", fake_get_model_response)
    monkeypatch.setattr(session, "_execute_tool", fake_execute_tool)

    asyncio.run(session._run_loop())

    assert executed == [("run_shell", {"command": "python --version"})]
    assert any(event.get("type") == "tool_loop_status" for event in session.events)
    assert session._verification_pending is False


def test_long_context_bundle_combines_memory_symbols_and_matches(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import long_context

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    monkeypatch.setattr(long_context.context_index.settings, "agent_workdir", str(workspace))
    monkeypatch.setattr(long_context.vector_index.settings, "agent_workdir", str(workspace))
    monkeypatch.setattr(long_context.vector_index.settings, "history_db_path", str(tmp_path / "sessions.db"))
    (workspace / "auth_service.py").write_text(
        "def login_user(email, password):\n    return create_session_token(email)\n\ndef create_session_token(email):\n    return email + '-token'\n",
        encoding="utf-8",
    )
    (workspace / "README.md").write_text("# Auth project\nHandles login and session token behavior.\n", encoding="utf-8")
    long_context.vector_index.rebuild_index()

    bundle = long_context.build_bundle("fix login session token", budget=5000)
    context = str(bundle["context"])

    assert "Project memory" in context
    assert "Symbol matches" in context
    assert "Semantic matches" in context
    assert "Exact keyword chunks" in context
    assert any(item["symbol"] == "login_user" for item in bundle["symbols"])
    assert "auth_service.py" in context


def test_context_bundle_route_and_agent_tool(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from fastapi.testclient import TestClient
    from app.agent import tools
    from app.main import app
    from app import long_context

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    monkeypatch.setattr(long_context.context_index.settings, "agent_workdir", str(workspace))
    monkeypatch.setattr(long_context.vector_index.settings, "agent_workdir", str(workspace))
    monkeypatch.setattr(long_context.vector_index.settings, "history_db_path", str(tmp_path / "sessions.db"))
    (workspace / "long_context_auth.py").write_text("def login_user():\n    return 'session token'\n", encoding="utf-8")

    result = TestClient(app).get("/api/context/bundle", params={"query": "login session token", "budget": 5000}).json()

    assert result["query"] == "login session token"
    assert "context" in result
    assert "login" in result["context"].lower()
    assert "long_context_bundle" in tools.READ_ONLY_TOOLS
    assert "Project memory" in tools.long_context_bundle("login session token")


def test_agent_injects_long_context_for_workspace_questions(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.agent.loop import AgentSession
    from app.agent import loop

    monkeypatch.setattr(loop.tools, "long_context_bundle", lambda query: "Project memory\n- app.py important")
    monkeypatch.setattr(loop.context_index, "automatic_context", lambda query: "exact context")
    monkeypatch.setattr(loop.context_index, "automatic_summary_context", lambda query: "summary context")
    monkeypatch.setattr(loop.context_index, "project_map_context", lambda query: "old project map")
    monkeypatch.setattr(loop.context_index, "conversation_memory", lambda events, query: "")

    packet = AgentSession(title="long context agent test")._with_file_context("explain this project structure", [])

    assert "Fused long-context bundle" in packet
    assert "Project memory" in packet
    assert "old project map" not in packet


def test_mcp_exposes_long_context_bundle(test_env: dict[str, str]) -> None:
    proc = subprocess.run(
        [sys.executable, "mcp_server.py"],
        input='{"jsonrpc":"2.0","id":1,"method":"tools/list"}\n',
        cwd=BACKEND_DIR,
        env=test_env,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    tool_names = {tool["name"] for tool in payload["result"]["tools"]}
    assert "long_context_bundle" in tool_names


def test_sql_reporting_schema_query_history_and_blocking(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import sqlite3
    from app import reports

    db_path = tmp_path / "sessions.db"
    monkeypatch.setattr(reports.settings, "history_db_path", str(db_path))
    with sqlite3.connect(db_path) as conn:
        conn.execute("CREATE TABLE sessions (id TEXT, title TEXT, owner_id TEXT)")
        conn.executemany(
            "INSERT INTO sessions (id, title, owner_id) VALUES (?, ?, ?)",
            [("1", "First demo", "local"), ("2", "Faculty report", "local")],
        )
        conn.commit()

    schema = reports.schema_overview()
    table_names = {table["name"] for table in schema["tables"]}
    assert "sessions" in table_names

    result = reports.run_report_query("SELECT title FROM sessions ORDER BY id", title="Session titles", limit=10)
    assert result["ok"] is True
    assert result["columns"] == ["title"]
    assert result["rows"] == [{"title": "First demo"}, {"title": "Faculty report"}]
    assert result["saved"] is True
    assert reports.report_history()[0]["title"] == "Session titles"

    with pytest.raises(ValueError):
        reports.run_report_query("DELETE FROM sessions")


def test_sql_reporting_routes(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import sqlite3
    from fastapi.testclient import TestClient
    from app import reports
    from app.main import app

    db_path = tmp_path / "sessions.db"
    monkeypatch.setattr(reports.settings, "history_db_path", str(db_path))
    with sqlite3.connect(db_path) as conn:
        conn.execute("CREATE TABLE sessions (id TEXT, title TEXT)")
        conn.execute("INSERT INTO sessions (id, title) VALUES ('1', 'Route report')")
        conn.commit()

    client = TestClient(app)
    schema = client.get("/api/reports/schema").json()
    assert any(table["name"] == "sessions" for table in schema["tables"])

    query = client.post("/api/reports/query", json={"sql": "SELECT title FROM sessions", "limit": 5}).json()
    assert query["row_count"] == 1
    assert query["rows"][0]["title"] == "Route report"

    blocked = client.post("/api/reports/query", json={"sql": "DROP TABLE sessions"})
    assert blocked.status_code == 400

    history = client.get("/api/reports/history").json()
    assert history["reports"][0]["row_count"] == 1


def test_dependency_scan_and_plan(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import dependencies

    root = tmp_path / "repo"
    frontend = root / "frontend"
    backend = root / "backend"
    frontend.mkdir(parents=True)
    backend.mkdir()
    (frontend / "package.json").write_text(
        json.dumps({"scripts": {"build": "vite build"}, "dependencies": {"react": "latest"}}),
        encoding="utf-8",
    )
    (backend / "requirements.txt").write_text("fastapi\n", encoding="utf-8")
    monkeypatch.setattr(dependencies, "repo_root", lambda: root)

    scan = dependencies.scan_projects()
    assert scan["summary"]
    assert any(project["path"] == "frontend" and "npm" in project["managers"] for project in scan["projects"])
    assert any(project["path"] == "backend" and "pip" in project["managers"] for project in scan["projects"])

    plan = dependencies.plan_dependencies("make a CRM dashboard with charts", target="frontend")
    packages = {item["package"] for item in plan["suggestions"]}
    assert "recharts" in packages
    assert any(command.startswith("cd frontend && npm install") for command in plan["commands"])
    assert "Ask approval" in " ".join(plan["workflow"])


def test_dependency_install_requires_approval_and_blocks_unsafe_package(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import dependencies

    root = tmp_path / "repo"
    frontend = root / "frontend"
    frontend.mkdir(parents=True)
    (frontend / "package.json").write_text("{}", encoding="utf-8")
    monkeypatch.setattr(dependencies, "repo_root", lambda: root)

    preview = dependencies.install_dependencies("npm", ["recharts"], "frontend", approve=False)
    assert preview["ran"] is False
    assert "Approval required" in preview["message"]
    assert preview["command"] == "npm install recharts"

    with pytest.raises(ValueError):
        dependencies.install_dependencies("npm", ["--force"], "frontend", approve=True)


def test_dependency_install_updates_requirements_after_success(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import dependencies

    root = tmp_path / "repo"
    backend = root / "backend"
    backend.mkdir(parents=True)
    (backend / "requirements.txt").write_text("fastapi\n", encoding="utf-8")
    monkeypatch.setattr(dependencies, "repo_root", lambda: root)

    class FakeProc:
        returncode = 0
        stdout = "installed"
        stderr = ""

    monkeypatch.setattr(dependencies.subprocess, "run", lambda *args, **kwargs: FakeProc())
    result = dependencies.install_dependencies("pip", ["pypdf"], "backend", approve=True, update_manifest=True)

    assert result["ok"] is True
    assert "pypdf" in (backend / "requirements.txt").read_text(encoding="utf-8")
    assert result["manifest_updates"] == ["Added pypdf to requirements.txt"]


def test_dependency_routes(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from fastapi.testclient import TestClient
    from app import dependencies
    from app.main import app

    root = tmp_path / "repo"
    frontend = root / "frontend"
    frontend.mkdir(parents=True)
    (frontend / "package.json").write_text(json.dumps({"dependencies": {}}), encoding="utf-8")
    monkeypatch.setattr(dependencies, "repo_root", lambda: root)

    client = TestClient(app)
    scan = client.get("/api/dependencies/scan").json()
    assert scan["projects"][0]["path"] == "frontend"

    plan = client.post("/api/dependencies/plan", json={"prompt": "make crm dashboard charts", "target": "frontend"}).json()
    assert any(item["package"] == "recharts" for item in plan["suggestions"])

    preview = client.post(
        "/api/dependencies/install",
        json={"manager": "npm", "packages": ["recharts"], "project_path": "frontend", "approve": False},
    ).json()
    assert preview["ran"] is False

    blocked = client.post(
        "/api/dependencies/install",
        json={"manager": "npm", "packages": ["--unsafe"], "project_path": "frontend", "approve": True},
    )
    assert blocked.status_code == 400


def test_agent_exposes_dependency_planning_tools(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.agent import tools

    monkeypatch.setattr(tools.dependencies, "scan_projects", lambda: {"projects": [], "summary": "none"})
    monkeypatch.setattr(tools.dependencies, "plan_dependencies", lambda prompt, target="auto": {"prompt": prompt, "target": target, "suggestions": []})

    assert "dependency_scan" in tools.READ_ONLY_TOOLS
    assert "dependency_plan" in tools.READ_ONLY_TOOLS
    assert "dependency_scan" in {item["function"]["name"] for item in tools.TOOL_SCHEMAS}
    assert "none" in tools.dependency_scan()
    assert "crm" in tools.dependency_plan("crm dashboard")


def test_mcp_exposes_dependency_tools(test_env: dict[str, str]) -> None:
    proc = subprocess.run(
        [sys.executable, "mcp_server.py"],
        input='{"jsonrpc":"2.0","id":1,"method":"tools/list"}\n',
        cwd=BACKEND_DIR,
        env=test_env,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    tool_names = {tool["name"] for tool in payload["result"]["tools"]}
    assert "dependency_scan" in tool_names
    assert "dependency_plan" in tool_names


def test_cli_prd_build_prompt_contains_required_workflow() -> None:
    import cli

    prompt = cli.build_prd_build_prompt("uploads/openbazaar.txt", "Marketplace PRD with buyer and seller listings")

    assert "PRD workspace path: uploads/openbazaar.txt" in prompt
    assert "Summarize the PRD requirements" in prompt
    assert "buyer/seller flows" in prompt
    assert "20-prompt build sequence" in prompt


def test_cli_prepare_prd_context_stages_absolute_text_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import cli

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = tmp_path / "OpenBazaar_PRD.txt"
    source.write_text("OpenBazaar marketplace requirements", encoding="utf-8")
    monkeypatch.setattr(cli.settings, "agent_workdir", str(workspace))

    rel_path, text = cli.prepare_prd_context(str(source))

    assert rel_path.startswith("uploads/cli-prd-")
    assert rel_path.endswith("OpenBazaar_PRD.txt")
    assert text == "OpenBazaar marketplace requirements"
    assert (workspace / rel_path).exists()


def test_cli_help_includes_prd_build(test_env: dict[str, str]) -> None:
    proc = subprocess.run(
        [sys.executable, "cli.py", "--help"],
        cwd=BACKEND_DIR,
        env=test_env,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert proc.returncode == 0, proc.stderr
    assert "prd-build" in proc.stdout



def test_enterprise_prd_mode_routes_complex_provider_prompts() -> None:
    from app.routes.tasks import build_plan

    plan = build_plan("Build this complex PRD with payment gateway, real OTP, courier workflow, cloud deployment, and multi-service backend business logic")

    assert plan.mode == "enterprise-prd-build-mode"
    assert "adapter interfaces" in plan.stack
    assert "enterprise_prd_project/adapters/payment_gateway.py" in plan.file_plan
    assert "enterprise PRD" in plan.workflow_summary
    assert any("provider contracts" in note for note in plan.notes)


def test_enterprise_prd_mode_scaffold_verifies(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.routes import tasks

    monkeypatch.setattr(tasks.settings, "agent_workdir", str(tmp_path))
    plan = tasks.build_plan("Create an enterprise PRD system with payment, OTP, courier, and deployment integrations")
    steps: list[tasks.TaskRunStep] = []
    created = tasks._write_suggested_files(plan, overwrite=True, steps=steps)
    ok = tasks._verify_created_files(plan, created, steps)

    assert ok is True
    assert "enterprise_prd_project/INTEGRATIONS.md" in created
    assert "PaymentGateway" in (tmp_path / "enterprise_prd_project" / "adapters" / "payment_gateway.py").read_text(encoding="utf-8")
    assert "OtpProvider" in (tmp_path / "enterprise_prd_project" / "adapters" / "otp_provider.py").read_text(encoding="utf-8")
    assert "CourierProvider" in (tmp_path / "enterprise_prd_project" / "adapters" / "courier_provider.py").read_text(encoding="utf-8")

def test_openbazaar_marketplace_prompt_uses_staged_roadmap_by_default(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.routes import tasks

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    monkeypatch.setattr(tasks.tools.settings, "agent_workdir", str(workspace))

    plan = tasks.build_plan("make the full OpenBazaar marketplace project with buyer dashboard seller dashboard auction cod admin")

    assert plan.mode == "openbazaar-20-step-prd-build-roadmap"
    assert plan.suggested_files == []
    assert len(plan.steps) == 20
    assert any("Step 1" in step for step in plan.steps)
    assert any("Step 20" in step for step in plan.steps)


def test_prd_build_prompt_routes_openbazaar_to_staged_roadmap() -> None:
    import cli
    from app.routes import tasks

    prd_text = """
    OpenBazaar marketplace PRD. Users include Guest, Buyer, Seller, and Admin.
    Features include Cash on Delivery, OTP verification, product listings, live auctions,
    bids, seller dashboard, buyer dashboard, orders, and admin moderation.
    """
    prompt = cli.build_prd_build_prompt("uploads/openbazaar-prd.txt", prd_text)
    plan = tasks.build_plan(prompt)

    assert plan.mode == "openbazaar-20-step-prd-build-roadmap"
    assert plan.suggested_files == []
    assert len(plan.steps) == 20








def test_openbazaar_summary_prompt_does_not_generate_files() -> None:
    from app.routes import tasks

    plan = tasks.build_plan(
        "Read this OpenBazaar PRD and summarize the roles, pages, database, "
        "security, and architecture requirements. Do not create files yet."
    )

    assert plan.mode == "openbazaar-prd-analysis"
    assert plan.suggested_files == []
    assert "No files are written" in " ".join(plan.file_plan)
    assert any("Cash on Delivery" in step for step in plan.steps)


def test_openbazaar_roadmap_prompt_does_not_generate_files() -> None:
    from app.routes import tasks

    plan = tasks.build_plan("Create a phase-by-phase roadmap for OpenBazaar. Do not build yet.")

    assert plan.mode == "openbazaar-prd-roadmap"
    assert plan.suggested_files == []
    assert any("Phase 1" in step for step in plan.steps)


def test_openbazaar_explicit_build_uses_staged_roadmap_by_default() -> None:
    from app.routes import tasks

    plan = tasks.build_plan(
        "Build the full OpenBazaar project with frontend backend database security auction and COD."
    )

    assert plan.mode == "openbazaar-20-step-prd-build-roadmap"
    assert plan.suggested_files == []
    assert len(plan.steps) == 20


def test_openbazaar_numbered_step_generates_only_that_step_files() -> None:
    from app.routes import tasks

    plan = tasks.build_plan("OpenBazaar step 4 create the application shell")

    assert plan.mode == "openbazaar-staged-prd-build-step-04"
    assert [item["path"] for item in plan.suggested_files] == ["openbazaar_marketplace/index.html"]
    assert "step 4 of 20" in " ".join(plan.notes).lower()


def test_openbazaar_later_step_does_not_rewrite_files() -> None:
    from app.routes import tasks

    plan = tasks.build_plan("OpenBazaar step 16 verify Cash on Delivery workflow")

    assert plan.mode == "openbazaar-staged-prd-build-step-16"
    assert plan.suggested_files == []
    assert any("smoke_test" in command for command in plan.verify_commands)


def test_openbazaar_emergency_full_build_is_explicit() -> None:
    from app.routes import tasks

    plan = tasks.build_plan("OpenBazaar emergency full generate everything now")

    assert plan.mode == "openbazaar-full-project-generator"
    assert any(item["path"] == "openbazaar_marketplace/index.html" for item in plan.suggested_files)




def test_toy_os_build_mode_generates_kernel_scaffold() -> None:
    from app.routes import tasks

    plan = tasks.build_plan("Make a toy operating system that boots in QEMU and prints Welcome to SHAMSU OS")

    paths = {item["path"] for item in plan.suggested_files}
    assert plan.mode == "toy-os-build-mode"
    assert "shamsu_os/boot.asm" in paths
    assert "shamsu_os/kernel.c" in paths
    assert "shamsu_os/linker.ld" in paths
    assert "shamsu_os/Makefile" in paths
    assert any("qemu-system-i386" in command for command in plan.verify_commands)
    assert any("not a production OS" in note for note in plan.notes)


def test_toy_os_run_writes_and_verifies_files(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import asyncio
    from app.agent import tools
    from app.routes import tasks

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    monkeypatch.setattr(tools.settings, "agent_workdir", str(workspace))
    monkeypatch.setattr(tasks.settings, "agent_workdir", str(workspace))

    result = asyncio.run(tasks.run_task(tasks.TaskRunRequest(prompt="Create a toy OS kernel for QEMU", preview=False)))

    assert result.ok is True
    assert "shamsu_os/kernel.c" in result.created_files
    assert (workspace / "shamsu_os" / "kernel.c").exists()
    assert "Welcome to SHAMSU OS" in (workspace / "shamsu_os" / "kernel.c").read_text(encoding="utf-8")
    assert any(step.name == "verify" and step.status == "ok" for step in result.steps)


def test_openbazaar_web_typo_prompt_routes_to_staged_build() -> None:
    from app.routes import tasks

    plan = tasks.build_plan("OpenBazzar step 4 create index.html from the PRD")

    assert plan.mode == "openbazaar-staged-prd-build-step-04"
    assert [item["path"] for item in plan.suggested_files] == ["openbazaar_marketplace/index.html"]


def test_openbazaar_spaced_name_prompt_routes_to_staged_build() -> None:
    from app.routes import tasks

    plan = tasks.build_plan("Open Bazaar step 5 add seed data")

    assert plan.mode == "openbazaar-staged-prd-build-step-05"
    assert [item["path"] for item in plan.suggested_files] == ["openbazaar_marketplace/src/data.js"]


def test_tool_workflow_brief_guides_build_bugfix_and_large_file_paths() -> None:
    from app.agent.loop import _tool_workflow_brief
    from app import query_policy

    build_policy = query_policy.classify_query("build a CRM project", [])
    build_brief = _tool_workflow_brief("build a CRM project", build_policy, [])
    assert "Build path" in build_brief
    assert "file plan" in build_brief
    assert "verify" in build_brief.lower()

    bug_policy = query_policy.classify_query("fix the login traceback in this project", [])
    bug_brief = _tool_workflow_brief("fix the login traceback in this project", bug_policy, [])
    assert "Debug path" in bug_brief
    assert "read only relevant ranges" in bug_brief
    assert "smallest exact block" in bug_brief

    large_policy = query_policy.classify_query("fix a bug in a 100000 line large file", [])
    large_brief = _tool_workflow_brief("fix a bug in a 100000 line large file", large_policy, [])
    assert "Large-project path" in large_brief or "Debug path" in large_brief
    assert "read_file_range" in large_brief or "relevant ranges" in large_brief


def test_tool_workflow_brief_stays_off_for_general_teaching_questions() -> None:
    from app.agent.loop import _tool_workflow_brief
    from app import query_policy

    policy = query_policy.classify_query("explain software engineering simply", [])

    assert _tool_workflow_brief("explain software engineering simply", policy, []) == ""


def test_agent_context_includes_claude_like_workflow_for_build_requests(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.agent.loop import AgentSession
    from app.agent import loop

    monkeypatch.setattr(loop.tools, "long_context_bundle", lambda query: "Project memory")
    monkeypatch.setattr(loop.context_index, "automatic_context", lambda query: "")
    monkeypatch.setattr(loop.context_index, "automatic_summary_context", lambda query: "")
    monkeypatch.setattr(loop.context_index, "project_map_context", lambda query: "")
    monkeypatch.setattr(loop.context_index, "conversation_memory", lambda events, query: "")

    packet = AgentSession(title="workflow brief test")._with_file_context("build a student management project", [])

    assert "Claude-like task workflow" in packet
    assert "Build path" in packet
    assert "User request: build a student management project" in packet


def test_project_understanding_report_guides_large_project_work(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import context_index

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    monkeypatch.setattr(context_index.settings, "agent_workdir", str(workspace))
    (workspace / "package.json").write_text('{"scripts":{"build":"vite build","test":"vitest"},"dependencies":{"react":"latest"}}', encoding="utf-8")
    (workspace / "src").mkdir()
    (workspace / "src" / "authService.ts").write_text("export function loginUser(email:string){ return email + '-token'; }\n", encoding="utf-8")
    (workspace / "src" / "main.tsx").write_text("import { loginUser } from './authService';\nconsole.log(loginUser('a'));\n", encoding="utf-8")

    report = context_index.project_understanding("fix login bug in auth service", limit=5)

    assert report["task_profile"]["kind"] == "bugfix"
    assert "react" in report["architecture"]["frameworks"]
    likely_paths = {item["path"] for item in report["likely_files"]}
    assert "src/authService.ts" in likely_paths
    assert any(item["path"] == "src/authService.ts" for item in report["read_plan"])
    assert "npm run build" in report["verification_commands"]


def test_project_understanding_context_is_compact_and_actionable(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import context_index

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    monkeypatch.setattr(context_index.settings, "agent_workdir", str(workspace))
    (workspace / "app.py").write_text("def create_order():\n    return 'ok'\n", encoding="utf-8")

    context = context_index.project_understanding_context("fix order bug")

    assert "Large-project understanding" in context
    assert "Task profile: bugfix" in context
    assert "Suggested read ranges" in context
    assert "app.py" in context


def test_agent_tool_exposes_project_understanding() -> None:
    from app.agent import tools

    tool_names = {schema["function"]["name"] for schema in tools.TOOL_SCHEMAS}

    assert "project_understanding" in tool_names
    assert "project_understanding" in tools.READ_ONLY_TOOLS
    assert "architecture" in tools.project_understanding("explain this project")


def test_context_understand_route_returns_read_plan(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from fastapi.testclient import TestClient
    from app import context_index
    from app.main import app

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    monkeypatch.setattr(context_index.settings, "agent_workdir", str(workspace))
    (workspace / "main.py").write_text("def main():\n    return 'hello'\n", encoding="utf-8")

    result = TestClient(app).get("/api/context/understand", params={"query": "fix main bug"}).json()

    assert result["task_profile"]["kind"] == "bugfix"
    assert result["read_plan"]
    assert result["likely_files"][0]["path"] == "main.py"


def test_long_context_bundle_includes_project_understanding(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import long_context

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    monkeypatch.setattr(long_context.context_index.settings, "agent_workdir", str(workspace))
    monkeypatch.setattr(long_context.vector_index.settings, "agent_workdir", str(workspace))
    monkeypatch.setattr(long_context.vector_index.settings, "history_db_path", str(tmp_path / "sessions.db"))
    (workspace / "orders.py").write_text("def create_order():\n    return 'pending'\n", encoding="utf-8")

    bundle = long_context.build_bundle("fix order creation bug", budget=5000, fast=True)

    assert "understanding" in bundle
    assert bundle["understanding"]["task_profile"]["kind"] == "bugfix"
    assert "Project understanding" in bundle["context"]
    assert "orders.py" in bundle["context"]


def test_conversation_memory_snapshot_groups_relevant_recent_and_files() -> None:
    from app import context_index

    events = [
        {"type": "user_message", "content": "build login dashboard"},
        {"type": "tool_call", "name": "read_file", "args": {"path": "auth.py"}},
        {"type": "approval_request", "name": "write_file", "path": "auth.py"},
        {"type": "approval_resolved", "approved": True},
        {"type": "files_changed", "paths": ["auth.py"]},
        {"type": "tool_result", "name": "run_shell", "ok": False, "preview": "login test failed"},
        {"type": "error", "message": "Login token mismatch"},
        {"type": "assistant_message", "content": "I will repair the login token check."},
    ]

    snapshot = context_index.conversation_memory_snapshot(events, "fix login token", budget=1200)

    assert "auth.py" in snapshot["files_touched"]
    assert any("login" in item.lower() for item in snapshot["relevant"] + snapshot["recent"])
    assert any("Tool result: run_shell failed" in item for item in snapshot["recent"] + snapshot["relevant"])


def test_conversation_memory_formats_smooth_sections_and_dedupes() -> None:
    from app import context_index

    events = [
        {"type": "user_message", "content": "create OpenBazaar shell"},
        {"type": "user_message", "content": "create OpenBazaar shell"},
        {"type": "approval_request", "name": "write_file", "path": "openbazaar_marketplace/index.html"},
        {"type": "approval_resolved", "approved": True},
        {"type": "files_changed", "paths": ["openbazaar_marketplace/index.html"]},
        {"type": "tool_result", "name": "write_file", "ok": True, "preview": "Wrote index.html"},
        {"type": "assistant_message", "content": "Created the shell."},
    ]

    memory = context_index.conversation_memory(events, "continue OpenBazaar index work", budget=1200)

    assert "Relevant earlier context" in memory
    assert "Files touched this session" in memory
    assert "Recent continuity" in memory
    assert memory.count("User asked: create OpenBazaar shell") == 1
    assert "openbazaar_marketplace/index.html" in memory


def test_conversation_memory_respects_small_budget() -> None:
    from app import context_index

    events = [
        {"type": "user_message", "content": f"important previous prompt {index} about login and dashboard"}
        for index in range(20)
    ]
    events.extend([
        {"type": "approval_request", "name": "write_file", "path": "dashboard.py"},
        {"type": "files_changed", "paths": ["dashboard.py"]},
    ])

    memory = context_index.conversation_memory(events, "login dashboard", budget=260)

    assert len(memory) <= 310
    assert "conversation memory truncated" in memory or "dashboard" in memory
